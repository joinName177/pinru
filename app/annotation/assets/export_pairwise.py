#!/usr/bin/env python3
import argparse
import json
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "User Prompt", "初始环境快照", "A-SessionID", "A-产物快照", "B-SessionID", "B-产物快照",
    "任务类型", "任务难度", "Harness", "Harness 版本", "操作系统", "环境可复现等级",
    "A-轨迹文件", "A-运行录屏", "B-轨迹文件", "B-运行录屏", "GSB 结论", "GSB 理由",
    "备注", "提交人", "提交时间", "项目", "题目", "待补材料",
]
CONCLUSIONS = {"A_better": "A 更好", "B_better": "B 更好", "same": "Same"}


def current_review(case):
    reviews = (case.get("pairwise") or {}).get("reviews") or []
    return reviews[-1] if reviews else {}


def row(case, payload):
    pairwise = case.get("pairwise") or {}
    run_a = pairwise.get("runA") or {}
    run_b = pairwise.get("runB") or {}
    review = current_review(case)
    return [
        pairwise.get("prompt", ""), case.get("snapshotUrl", ""),
        run_a.get("sessionId", ""), run_a.get("deliverableUrl", ""),
        run_b.get("sessionId", ""), run_b.get("deliverableUrl", ""),
        case.get("taskType", ""), case.get("promptDifficulty", ""),
        pairwise.get("harness", ""), pairwise.get("harnessVersion", ""), pairwise.get("os", ""),
        pairwise.get("environment", ""), run_a.get("tracePath", ""),
        run_a.get("videoUrl") or run_a.get("videoPath", ""), run_b.get("tracePath", ""),
        run_b.get("videoUrl") or run_b.get("videoPath", ""), CONCLUSIONS.get(review.get("conclusion"), ""),
        review.get("reason", ""), pairwise.get("notes", ""), payload.get("submitter", ""),
        payload.get("submittedAt", ""), payload.get("projectName", ""), case.get("taskName", ""),
        "；".join(payload.get("issues") or []),
    ]


def export(payload, output):
    output.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pair-wise GSB"
    sheet.append(HEADERS)
    for case in payload.get("cases") or []:
        sheet.append(row(case, payload))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [36, 48, 24, 48, 24, 48, 18, 14, 16, 16, 18, 24, 42, 42, 42, 42, 14, 72, 32, 16, 20, 22, 24, 60]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook_path = (output / "pairwise-gsb.xlsx").resolve()
    workbook.save(workbook_path)

    issues = payload.get("issues") or []
    report_path = (output / "pairwise-report.md").resolve()
    lines = ["# Pair-wise GSB 导出报告", "", f"- 题目数：{len(payload.get('cases') or [])}", f"- 待补项：{len(issues)}"]
    if issues:
        lines.extend(["", "## 待补材料", ""] + [f"- {item}" for item in issues])
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {
        "outputPath": str(workbook_path),
        "reportPath": str(report_path),
        "rows": len(payload.get("cases") or []),
        "issues": issues,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--draft", action="store_true")
    args = parser.parse_args()
    try:
        payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
        print(json.dumps(export(payload, Path(args.output)), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"outputPath": "", "reportPath": "", "rows": 0, "issues": [str(exc)]}, ensure_ascii=False))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
